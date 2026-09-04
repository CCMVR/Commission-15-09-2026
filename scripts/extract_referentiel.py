import json
import re
import unicodedata
import openpyxl

def normalize_key(s):
    if not s:
        return ""
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = re.sub(r"[^a-zA-Z0-9]+", " ", s).strip().lower()
    return re.sub(r"\s+", " ", s)

# Dictionnaire de correspondance canonique
CANONICAL_COMMUNES = {
    "bas en basset": "Bas-en-Basset",
    "beauzac": "Beauzac",
    "boisset": "Boisset",
    "la chapelle d aurec": "La Chapelle-d'Aurec",
    "les villettes": "Les Villettes",
    "malvalette": "Malvalette",
    "monistrol sur loire": "Monistrol-sur-Loire",
    "saint andre de chalencon": "Saint-André-de-Chalencon",
    "saint pal de chalencon": "Saint-Pal-de-Chalencon",
    "saint pal de mons": "Saint-Pal-de-Mons",
    "sainte sigolene": "Sainte-Sigolène",
    "solignac sous roche": "Solignac-sous-Roche",
    "tiranges": "Tiranges",
    "valprivas": "Valprivas",
}

def to_canonical_commune(name):
    if not name:
        return ""
    key = normalize_key(name)
    return CANONICAL_COMMUNES.get(key, name.strip())

# 1. Extraction des structures
wb_structures = openpyxl.load_workbook("Fiche structures.xlsx")
ws_structures = wb_structures.active

structures = []
for i, row in enumerate(ws_structures.iter_rows(values_only=True)):
    if i == 0:
        continue  # Header: Nom, Type, Ville
    nom, type_service, ville = row[0], row[1], row[2]
    if not nom or not ville:
        continue
    nom = str(nom).strip()
    type_service = str(type_service).strip() if type_service else "ALSH"
    commune_canon = to_canonical_commune(str(ville).strip())
    
    # Découpage des types de services
    # Exemple: "Crèche et ALSH et Ados" -> ["Crèche", "ALSH", "Ados"]
    services_raw = re.split(r"\s+et\s+|,\s*|/\s*", type_service)
    services = [s.strip() for s in services_raw if s.strip()]
    
    slug = re.sub(r"[^a-z0-9]+", "-", normalize_key(nom)).strip("-")
    structures.append({
        "id": slug,
        "name": nom,
        "commune": commune_canon,
        "services": services,
        "servicesLabel": type_service
    })

print(f"Total structures extraites: {len(structures)}")

# 2. Extraction des élus
wb_elus = openpyxl.load_workbook("Liste des élus.xlsx")
ws_elus = wb_elus.active

elus = []
for i, row in enumerate(ws_elus.iter_rows(values_only=True)):
    if not row or not row[0]:
        continue
    nom_brut = str(row[0]).strip()
    val2 = str(row[1]).strip() if row[1] else ""
    
    # Ligne 1 : Présidente
    if "DUPLAIN Jocelyne" in nom_brut:
        elus.append({
            "id": "duplain-jocelyne",
            "name": "DUPLAIN Jocelyne",
            "commune": "CC Marches du Velay Rochebaron",
            "role": "Présidente de la Communauté de communes"
        })
        continue
    
    # Ligne 2 : En-tête de section
    if "Elus membres" in nom_brut or "Communes" in val2:
        continue
    
    commune_canon = to_canonical_commune(val2)
    slug = re.sub(r"[^a-z0-9]+", "-", normalize_key(nom_brut)).strip("-")
    elus.append({
        "id": slug,
        "name": nom_brut,
        "commune": commune_canon,
        "role": "Membre de la commission"
    })

print(f"Total élus extraits: {len(elus)}")

# 3. Communes avec structures ordonnées
communes_avec_structures = sorted(list(set(s["commune"] for s in structures)))
print(f"Communes avec structures ({len(communes_avec_structures)}): {communes_avec_structures}")

# Structure par commune
structures_par_commune = {}
for c in communes_avec_structures:
    structures_par_commune[c] = [s for s in structures if s["commune"] == c]

# Génération de js/data.js
js_content = f"""/**
 * Référentiel unifié des données de la Commission Enfance Jeunesse
 * Généré automatiquement à partir de Fiche structures.xlsx et Liste des élus.xlsx
 */

export const CANONICAL_COMMUNES = {json.dumps(CANONICAL_COMMUNES, ensure_ascii=False, indent=2)};

export function normalizeCommune(name) {{
    if (!name) return "";
    return name
        .normalize("NFD")
        .replace(/[\u0300-\u036f]/g, "")
        .replace(/[^a-zA-Z0-9]+/g, " ")
        .trim()
        .toLowerCase();
}}

export function isCommuneBlocked(eluCommune, structureCommune) {{
    if (!eluCommune || !structureCommune) return false;
    const n1 = normalizeCommune(eluCommune);
    const n2 = normalizeCommune(structureCommune);
    return n1 === n2;
}}

export const COMMUNES_AVEC_STRUCTURES = {json.dumps(communes_avec_structures, ensure_ascii=False, indent=2)};

export const STRUCTURES = {json.dumps(structures, ensure_ascii=False, indent=2)};

export const ELUS = {json.dumps(elus, ensure_ascii=False, indent=2)};
"""

with open("js/data.js", "w", encoding="utf-8") as f:
    f.write(js_content)

print("js/data.js généré avec succès.")
